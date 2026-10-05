from typing import Optional, List, Dict
from uuid import UUID
from io import BytesIO
from datetime import datetime
from fastapi import APIRouter, Depends, Query, UploadFile, File, Request
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from app.core.exceptions import NotFoundError, BadRequestError, ConflictError
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func
from sqlalchemy.orm import selectinload
from app.database import get_db
from app.api.deps import get_current_user, RequirePermission
from app.schemas.job_position import JobPositionCreate, JobPositionUpdate, JobPositionResponse
from app.models.user import User
from app.models.job_position import JobPosition
from app.models.employee import Employee
from app.models.workplace import Workplace
from app.services.audit_service import create_audit_log

router = APIRouter()


async def _build_response(db: AsyncSession, jp: JobPosition, current_headcount: int) -> JobPositionResponse:
    wp_name = None
    if jp.workplace_id:
        if jp.workplace:
            wp_name = jp.workplace.name
        else:
            wp = await db.get(Workplace, jp.workplace_id)
            if wp:
                wp_name = wp.name

    return JobPositionResponse(
        id=str(jp.id),
        name=jp.name,
        code=jp.code,
        description=jp.description,
        professional_level=jp.professional_level,
        required_qualification_level=jp.required_qualification_level,
        required_specialty=jp.required_specialty,
        max_headcount=jp.max_headcount,
        workplace_id=str(jp.workplace_id) if jp.workplace_id else None,
        workplace_name=wp_name,
        is_active=jp.is_active,
        current_headcount=current_headcount,
        created_at=jp.created_at,
        updated_at=jp.updated_at,
    )


@router.get("", response_model=List[JobPositionResponse])
async def list_job_positions(
    active_only: bool = Query(True),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    query = select(JobPosition).options(selectinload(JobPosition.workplace))
    if active_only:
        query = query.where(JobPosition.is_active == True)
    query = query.order_by(JobPosition.professional_level, JobPosition.name)
    result = await db.execute(query)
    job_positions = result.scalars().all()

    headcounts: Dict[str, int] = {}
    count_result = await db.execute(
        select(Employee.job_position_id, func.count(Employee.id))
        .where(Employee.is_active == True, Employee.job_position_id.isnot(None))
        .group_by(Employee.job_position_id)
    )
    for jp_id, cnt in count_result:
        headcounts[str(jp_id)] = cnt

    return [
        await _build_response(db, jp, headcounts.get(str(jp.id), 0))
        for jp in job_positions
    ]


@router.get("/import/template")
async def download_job_positions_template(
    current_user: User = Depends(get_current_user),
):
    wb = Workbook()
    ws = wb.active
    ws.title = "ViTriViecLam"
    headers = [
        "Mã vị trí việc làm",
        "Tên vị trí việc làm",
        "Bậc nghề nghiệp",
        "Nơi làm việc",
        "Bằng cấp (trình độ)",
        "Chuyên môn",
        "Số lượng tối đa (Định biên)",
        "Mô tả",
    ]
    ws.append(headers)
    ws.append(["CNTT", "Chuyên viên Công nghệ thông tin", 1, "Trụ sở chính", "Đại học", "Công nghệ thông tin", 2, "Quản trị hệ thống IT"])
    ws.append(["KT", "Kế toán viên", 2, "Trụ sở chính", "Đại học", "Kế toán", 3, "Kế toán tổng hợp"])

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill

    for col in ws.columns:
        max_len = max(len(str(c.value or "")) for c in col) + 2
        ws.column_dimensions[col[0].column_letter].width = max(max_len, 15)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=mau_nhap_vitrivieclam.xlsx"},
    )


@router.get("/export")
async def export_job_positions_to_excel(
    active_only: bool = Query(False),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    query = select(JobPosition).options(selectinload(JobPosition.workplace))
    if active_only:
        query = query.where(JobPosition.is_active == True)
    query = query.order_by(JobPosition.professional_level, JobPosition.name)
    result = await db.execute(query)
    job_positions = result.scalars().all()

    headcounts: Dict[str, int] = {}
    count_result = await db.execute(
        select(Employee.job_position_id, func.count(Employee.id))
        .where(Employee.is_active == True, Employee.job_position_id.isnot(None))
        .group_by(Employee.job_position_id)
    )
    for jp_id, cnt in count_result:
        headcounts[str(jp_id)] = cnt

    wb = Workbook()
    ws = wb.active
    ws.title = "Vị trí việc làm"

    headers = [
        "Mã vị trí việc làm",
        "Tên vị trí việc làm",
        "Bậc nghề nghiệp",
        "Nơi làm việc",
        "Bằng cấp (trình độ)",
        "Chuyên môn",
        "Số lượng tối đa (Định biên)",
        "Số lượng hiện tại",
        "Trạng thái",
        "Mô tả",
    ]

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    for row_idx, jp in enumerate(job_positions, 2):
        row_data = [
            jp.code,
            jp.name,
            f"Bậc {jp.professional_level}",
            jp.workplace.name if jp.workplace else "",
            jp.required_qualification_level or "",
            jp.required_specialty or "",
            jp.max_headcount,
            headcounts.get(str(jp.id), 0),
            "Hoạt động" if jp.is_active else "Ngừng hoạt động",
            jp.description or "",
        ]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border
            if col_idx in (1, 3, 7, 8, 9):
                cell.alignment = Alignment(horizontal="center")

    for col in ws.columns:
        max_length = 0
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        adjusted_width = min(max(max_length + 3, 12), 50)
        ws.column_dimensions[col[0].column_letter].width = adjusted_width

    ws.auto_filter.ref = ws.dimensions

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=vitrivieclam_{timestamp}.xlsx"},
    )


@router.post("/import", status_code=200)
async def import_job_positions_from_excel(
    request: Request,
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
    perms: dict = Depends(RequirePermission('can_manage_settings')),
):
    if not file.filename or not file.filename.endswith(('.xlsx', '.xls')):
        raise BadRequestError(detail="Chỉ hỗ trợ file định dạng Excel (.xlsx, .xls)")

    content = await file.read()
    if not content:
        raise BadRequestError(detail="File Excel rỗng")

    wb = load_workbook(filename=BytesIO(content))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise BadRequestError(detail="File Excel rỗng")

    headers = [str(c).strip() if c is not None else "" for c in rows[0]]
    
    col_code = -1
    col_name = -1
    col_level = -1
    col_wp = -1
    col_qual = -1
    col_spec = -1
    col_max = -1
    col_desc = -1

    for idx, h in enumerate(headers):
        h_lower = h.lower()
        if "mã" in h_lower:
            col_code = idx
        elif "tên" in h_lower:
            col_name = idx
        elif "bậc" in h_lower:
            col_level = idx
        elif "nơi làm việc" in h_lower or "địa điểm" in h_lower or "trụ sở" in h_lower:
            col_wp = idx
        elif "bằng cấp" in h_lower or "trình độ" in h_lower:
            col_qual = idx
        elif "chuyên môn" in h_lower:
            col_spec = idx
        elif "tối đa" in h_lower or "định biên" in h_lower or "số lượng" in h_lower:
            col_max = idx
        elif "mô tả" in h_lower:
            col_desc = idx

    if col_code == -1 or col_name == -1:
        raise BadRequestError(detail="File Excel phải chứa ít nhất 2 cột: 'Mã vị trí việc làm' và 'Tên vị trí việc làm'")

    wp_cache: Dict[str, UUID] = {}
    wps_res = await db.execute(select(Workplace).where(Workplace.is_active == True))
    for w in wps_res.scalars().all():
        wp_cache[w.name.lower().strip()] = w.id
        if w.code:
            wp_cache[w.code.lower().strip()] = w.id

    created_count = 0
    updated_count = 0
    errors = []

    for row_idx, row in enumerate(rows[1:], start=2):
        if all(c is None or str(c).strip() == "" for c in row):
            continue

        code_val = str(row[col_code]).strip() if col_code < len(row) and row[col_code] else ""
        name_val = str(row[col_name]).strip() if col_name < len(row) and row[col_name] else ""

        if not code_val or not name_val:
            errors.append(f"Dòng {row_idx}: Mã và Tên vị trí việc làm không được để trống")
            continue

        prof_level = 1
        if col_level != -1 and col_level < len(row) and row[col_level] is not None:
            raw_lvl = str(row[col_level]).strip().lower().replace("bậc", "").strip()
            try:
                prof_level = int(raw_lvl)
                prof_level = max(1, min(5, prof_level))
            except ValueError:
                prof_level = 1

        wp_val = str(row[col_wp]).strip() if col_wp != -1 and col_wp < len(row) and row[col_wp] else None
        wp_id = wp_cache.get(wp_val.lower()) if wp_val else None

        qual_val = str(row[col_qual]).strip() if col_qual != -1 and col_qual < len(row) and row[col_qual] else None
        spec_val = str(row[col_spec]).strip() if col_spec != -1 and col_spec < len(row) and row[col_spec] else None
        
        max_hc = 1
        if col_max != -1 and col_max < len(row) and row[col_max] is not None:
            try:
                max_hc = int(row[col_max])
                max_hc = max(0, max_hc)
            except ValueError:
                max_hc = 1

        desc_val = str(row[col_desc]).strip() if col_desc != -1 and col_desc < len(row) and row[col_desc] else None

        existing_res = await db.execute(select(JobPosition).where(JobPosition.code == code_val))
        existing_jp = existing_res.scalar_one_or_none()

        if existing_jp:
            existing_jp.name = name_val
            existing_jp.professional_level = prof_level
            if wp_id is not None:
                existing_jp.workplace_id = wp_id
            if qual_val is not None:
                existing_jp.required_qualification_level = qual_val
            if spec_val is not None:
                existing_jp.required_specialty = spec_val
            existing_jp.max_headcount = max_hc
            if desc_val is not None:
                existing_jp.description = desc_val
            existing_jp.is_active = True
            updated_count += 1
        else:
            new_jp = JobPosition(
                code=code_val,
                name=name_val,
                professional_level=prof_level,
                workplace_id=wp_id,
                required_qualification_level=qual_val,
                required_specialty=spec_val,
                max_headcount=max_hc,
                description=desc_val,
                is_active=True,
            )
            db.add(new_jp)
            created_count += 1

    await db.commit()

    await create_audit_log(
        db, action="job_position.import", resource_type="job_position",
        user_id=str(current_user.id),
        new_values={"created": created_count, "updated": updated_count, "errors": len(errors)},
        request=request, success=True,
    )

    return {
        "imported_count": created_count + updated_count,
        "created_count": created_count,
        "updated_count": updated_count,
        "errors": errors,
    }


@router.get("/{jp_id}", response_model=JobPositionResponse)
async def get_job_position(
    jp_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    try:
        jp_uuid = UUID(jp_id)
    except ValueError:
        raise NotFoundError(detail="Không tìm thấy vị trí việc làm")
    result = await db.execute(
        select(JobPosition)
        .options(selectinload(JobPosition.workplace))
        .where(JobPosition.id == jp_uuid)
    )
    jp = result.scalar_one_or_none()
    if not jp:
        raise NotFoundError(detail="Không tìm thấy vị trí việc làm")

    emp_count_result = await db.execute(
        select(func.count(Employee.id))
        .where(Employee.job_position_id == jp_uuid, Employee.is_active == True)
    )
    emp_count = emp_count_result.scalar() or 0

    return await _build_response(db, jp, emp_count)


@router.post("", response_model=JobPositionResponse, status_code=201)
async def create_job_position(
    body: JobPositionCreate,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
    perms: dict = Depends(RequirePermission('can_manage_settings')),
):
    existing = await db.execute(select(JobPosition).where(JobPosition.code == body.code))
    if existing.scalar_one_or_none():
        raise ConflictError(detail="Mã vị trí việc làm đã tồn tại")

    data = body.model_dump()
    if data.get("workplace_id"):
        try:
            data["workplace_id"] = UUID(data["workplace_id"])
        except ValueError:
            data["workplace_id"] = None
    else:
        data["workplace_id"] = None

    jp = JobPosition(**data)
    db.add(jp)
    await db.flush()
    await db.refresh(jp)
    await create_audit_log(
        db, action="job_position.create", resource_type="job_position",
        resource_id=str(jp.id), user_id=str(current_user.id),
        new_values={"name": jp.name, "code": jp.code, "max_headcount": jp.max_headcount},
        success=True,
    )
    await db.commit()

    return await _build_response(db, jp, 0)


@router.put("/{jp_id}", response_model=JobPositionResponse)
async def update_job_position(
    jp_id: str,
    body: JobPositionUpdate,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
    perms: dict = Depends(RequirePermission('can_manage_settings')),
):
    try:
        jp_uuid = UUID(jp_id)
    except ValueError:
        raise NotFoundError(detail="Không tìm thấy vị trí việc làm")
    result = await db.execute(
        select(JobPosition)
        .options(selectinload(JobPosition.workplace))
        .where(JobPosition.id == jp_uuid)
    )
    jp = result.scalar_one_or_none()
    if not jp:
        raise NotFoundError(detail="Không tìm thấy vị trí việc làm")

    update_data = body.model_dump(exclude_unset=True)
    if "code" in update_data and update_data["code"] != jp.code:
        existing = await db.execute(select(JobPosition).where(JobPosition.code == update_data["code"]))
        if existing.scalar_one_or_none():
            raise ConflictError(detail="Mã vị trí việc làm đã tồn tại")

    if "workplace_id" in update_data:
        if update_data["workplace_id"]:
            try:
                update_data["workplace_id"] = UUID(update_data["workplace_id"])
            except ValueError:
                update_data["workplace_id"] = None
        else:
            update_data["workplace_id"] = None

    old_values = {k: getattr(jp, k, None) for k in update_data.keys()}
    for key, value in update_data.items():
        setattr(jp, key, value)

    await db.flush()
    await db.refresh(jp)

    emp_count_result = await db.execute(
        select(func.count(Employee.id))
        .where(Employee.job_position_id == jp_uuid, Employee.is_active == True)
    )
    emp_count = emp_count_result.scalar() or 0

    new_values = {k: (str(v) if isinstance(v, UUID) else v) for k, v in update_data.items()}
    await create_audit_log(
        db, action="job_position.update", resource_type="job_position",
        resource_id=str(jp.id), user_id=str(current_user.id),
        old_values=old_values, new_values=new_values, success=True,
    )
    await db.commit()

    return await _build_response(db, jp, emp_count)


@router.delete("/{jp_id}")
async def delete_job_position(
    jp_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
    perms: dict = Depends(RequirePermission('can_manage_settings')),
):
    try:
        jp_uuid = UUID(jp_id)
    except ValueError:
        raise NotFoundError(detail="Không tìm thấy vị trí việc làm")
    result = await db.execute(select(JobPosition).where(JobPosition.id == jp_uuid))
    jp = result.scalar_one_or_none()
    if not jp:
        raise NotFoundError(detail="Không tìm thấy vị trí việc làm")

    emp_count = await db.execute(
        select(func.count(Employee.id))
        .where(Employee.job_position_id == jp_uuid, Employee.is_active == True)
    )
    if emp_count.scalar() > 0:
        raise BadRequestError(detail="Không thể xoá vị trí việc làm đang có nhân sự")

    jp.is_active = False
    await create_audit_log(
        db, action="job_position.delete", resource_type="job_position",
        resource_id=jp_id, user_id=str(current_user.id),
        old_values={"is_active": True}, new_values={"is_active": False}, success=True,
    )
    await db.commit()
    return {"message": "Đã vô hiệu hóa vị trí việc làm"}


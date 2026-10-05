import logging
from io import BytesIO
from typing import List, Optional
from uuid import UUID
from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select
from sqlalchemy.orm import selectinload
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from app.database import get_db
from app.api.deps import get_current_user, RequirePermission, get_current_permissions, get_scope_filter
from app.models.employee import Employee
from app.models.user import User
from app.schemas.employee import EmployeeResponse
from app.core.constants import RESPONSE_MESSAGES
from app.core.exceptions import NotFoundError, ConflictError
from app.services.employee_account_service import auto_provision_user
from app.services.audit_service import create_audit_log

logger = logging.getLogger(__name__)

router = APIRouter()


class BatchEmployeeRequest(BaseModel):
    ids: List[str]


@router.post("/batch", response_model=List[EmployeeResponse])
async def batch_get_employees(
    body: BatchEmployeeRequest,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    uuid_ids = [UUID(eid) for eid in body.ids]
    result = await db.execute(
        select(Employee)
        .options(selectinload(Employee.department), selectinload(Employee.workplace), selectinload(Employee.position_rel))
        .where(Employee.id.in_(uuid_ids))
    )
    employees = result.scalars().all()
    emp_map = {str(emp.id): emp for emp in employees}

    user_result = await db.execute(
        select(User).where(User.employee_id.in_(uuid_ids))
    )
    user_map = {str(u.employee_id): u for u in user_result.scalars().all()}

    results = []
    for emp_id in body.ids:
        emp = emp_map.get(emp_id)
        if not emp:
            continue
        user = user_map.get(emp_id)
        has_account = user is not None
        results.append(EmployeeResponse(
            id=str(emp.id),
            department_id=str(emp.department_id),
            workplace_id=str(emp.workplace_id) if emp.workplace_id else None,
            position_id=str(emp.position_id) if emp.position_id else None,
            full_name=emp.full_name,
            position=emp.position,
            phone=emp.phone,
            email=emp.email,
            employee_code=emp.employee_code,
            hire_date=emp.hire_date,
            is_active=emp.is_active,
            status=emp.status,
            department_name=emp.department.name if emp.department else None,
            workplace_name=emp.workplace.name if emp.workplace else None,
            position_name=emp.position_rel.name if emp.position_rel else None,
            has_account=has_account,
            user_id=str(user.id) if user else None,
            username=user.username if user else None,
            is_super_admin=user.is_super_admin if user else False,
            created_at=emp.created_at,
            updated_at=emp.updated_at,
        ))
    return results


class BatchCreateAccountsRequest(BaseModel):
    employee_ids: List[str]


class BatchDeactivateRequest(BaseModel):
    employee_ids: List[str]
    status: str


@router.post("/batch/create-accounts")
async def batch_create_accounts(
    body: BatchCreateAccountsRequest,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(RequirePermission('can_manage_employees')),
):
    rows = []
    for emp_id in body.employee_ids:
        emp_uuid = UUID(emp_id)
        result = await db.execute(
            select(Employee).options(selectinload(Employee.department)).where(Employee.id == emp_uuid)
        )
        emp = result.scalar_one_or_none()
        if not emp:
            rows.append({"employee_code": "", "full_name": "", "username": "", "password": "", "status": f"Không tìm thấy nhân viên {emp_id}"})
            continue
        if not emp.is_active:
            rows.append({"employee_code": emp.employee_code or "", "full_name": emp.full_name, "username": "", "password": "", "status": "Nhân viên đã nghỉ việc"})
            continue
        try:
            user, generated_password = await auto_provision_user(db, emp, None)
            rows.append({
                "employee_code": emp.employee_code or "",
                "full_name": emp.full_name,
                "username": user.username,
                "password": generated_password,
                "status": "Thành công",
            })
        except ConflictError:
            rows.append({
                "employee_code": emp.employee_code or "",
                "full_name": emp.full_name,
                "username": "",
                "password": "",
                "status": "Đã có tài khoản",
            })
        except Exception as e:
            logger.exception(f"Lỗi tạo tài khoản cho {emp_id}")
            rows.append({
                "employee_code": emp.employee_code or "",
                "full_name": emp.full_name,
                "username": "",
                "password": "",
                "status": f"Lỗi: {str(e)[:200]}",
            })

    await db.flush()

    await create_audit_log(
        db, action="employee.batch_create_accounts", resource_type="employee",
        resource_id=",".join(body.employee_ids), user_id=str(current_user.id),
        new_values={
            "total": len(body.employee_ids),
            "success": sum(1 for r in rows if r["status"] == "Thành công"),
        },
        success=True,
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Taikhoan"
    headers = ["Mã NV", "Họ tên", "Tên đăng nhập", "Mật khẩu", "Trạng thái"]
    header_font = Font(bold=True)
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row_idx, row in enumerate(rows, 2):
        ws.cell(row=row_idx, column=1, value=row["employee_code"])
        ws.cell(row=row_idx, column=2, value=row["full_name"])
        ws.cell(row=row_idx, column=3, value=row["username"])
        ws.cell(row=row_idx, column=4, value=row["password"])
        ws.cell(row=row_idx, column=5, value=row["status"])

    for col in ws.columns:
        max_len = max(len(str(c.value or "")) for c in col) + 2
        ws.column_dimensions[col[0].column_letter].width = max(max_len, 12)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=taikhoan_nhanvien.xlsx"},
    )


@router.post("/batch/deactivate")
async def batch_deactivate(
    body: BatchDeactivateRequest,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(RequirePermission('can_manage_employees')),
):
    success_count = 0
    errors = []
    for emp_id in body.employee_ids:
        try:
            emp_uuid = UUID(emp_id)
            result = await db.execute(select(Employee).where(Employee.id == emp_uuid))
            emp = result.scalar_one_or_none()
            if not emp:
                errors.append({"id": emp_id, "reason": "Không tìm thấy"})
                continue
            if not emp.is_active:
                errors.append({"id": emp_id, "reason": "Đã nghỉ việc"})
                continue
            old_values = {"is_active": emp.is_active, "status": emp.status}
            emp.is_active = False
            emp.status = body.status

            user_result = await db.execute(select(User).where(User.employee_id == emp_id))
            user = user_result.scalar_one_or_none()
            if user:
                user.is_active = False

            await create_audit_log(
                db, action="employee.deactivate", resource_type="employee",
                resource_id=emp_id, user_id=str(current_user.id),
                old_values=old_values, new_values={"is_active": False, "status": body.status}, success=True,
            )
            success_count += 1
        except Exception as e:
            logger.exception(f"Lỗi deactivate {emp_id}")
            errors.append({"id": emp_id, "reason": str(e)[:200]})

    await db.flush()

    await create_audit_log(
        db, action="employee.batch_deactivate", resource_type="employee",
        resource_id=",".join(body.employee_ids), user_id=str(current_user.id),
        new_values={
            "total": len(body.employee_ids),
            "success": success_count,
            "errors": len(errors),
        },
        success=True,
    )

    return {
        "message": f"Đã ngưng {success_count} nhân viên" + (f", {len(errors)} lỗi" if errors else ""),
        "success_count": success_count,
        "errors": errors,
    }


class BatchHardDeleteRequest(BaseModel):
    employee_ids: List[str]


@router.post("/batch/hard-delete")
async def batch_hard_delete(
    body: BatchHardDeleteRequest,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(RequirePermission('can_manage_employees')),
):
    from sqlalchemy import update, delete
    from app.models.department import Department
    from app.models.refresh_token import RefreshToken

    success_count = 0
    errors = []
    for emp_id in body.employee_ids:
        try:
            emp_uuid = UUID(emp_id)
            result = await db.execute(select(Employee).where(Employee.id == emp_uuid))
            emp = result.scalar_one_or_none()
            if not emp:
                errors.append({"id": emp_id, "reason": "Không tìm thấy"})
                continue

            await db.execute(
                update(Department)
                .where(Department.head_id == emp.id)
                .values(head_id=None)
            )

            user_result = await db.execute(select(User).where(User.employee_id == emp.id))
            user = user_result.scalar_one_or_none()
            if user:
                await db.execute(delete(RefreshToken).where(RefreshToken.user_id == user.id))
                await db.delete(user)

            await create_audit_log(
                db, action="employee.hard_delete", resource_type="employee",
                resource_id=emp_id, user_id=str(current_user.id),
                old_values={"full_name": emp.full_name, "employee_code": emp.employee_code},
                new_values=None, success=True,
            )
            await db.delete(emp)
            success_count += 1
        except Exception as e:
            logger.exception(f"Lỗi hard delete {emp_id}")
            errors.append({"id": emp_id, "reason": str(e)[:200]})

    await db.flush()

    await create_audit_log(
        db, action="employee.batch_hard_delete", resource_type="employee",
        resource_id=",".join(body.employee_ids), user_id=str(current_user.id),
        new_values={
            "total": len(body.employee_ids),
            "success": success_count,
            "errors": len(errors),
        },
        success=True,
    )

    return {
        "message": f"Đã xóa vĩnh viễn {success_count} nhân viên" + (f", {len(errors)} lỗi" if errors else ""),
        "success_count": success_count,
        "errors": errors,
    }

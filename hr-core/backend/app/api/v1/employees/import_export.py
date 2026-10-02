import logging
from io import BytesIO
from datetime import datetime
from fastapi import APIRouter, Depends, UploadFile, File, Query, Request
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select
from sqlalchemy.orm import selectinload
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from app.database import get_db
from app.api.deps import get_current_user, RequirePermission, get_current_permissions, get_scope_filter
from app.models.user import User
from app.models.employee import Employee
from app.core.constants import RESPONSE_MESSAGES
from app.core.exceptions import BadRequestError
from app.core.limiter import limiter
from app.services.import_service import parse_excel, import_employees
from app.services.audit_service import create_audit_log

logger = logging.getLogger(__name__)

router = APIRouter()


@router.get("/import/template")
async def download_import_template(
    current_user: User = Depends(get_current_user),
):
    wb = Workbook()
    ws = wb.active
    ws.title = "Employees"
    headers = ["Mã NV", "Họ tên", "Phòng ban", "Chức vụ", "Nơi làm việc", "Ngày vào làm", "SĐT", "Email"]
    ws.append(headers)
    ws.append(["NV001", "Nguyễn Văn A", "OFFICE", "Nhân viên", "HQ", "01/01/2024", "0909123456", "a.nguyen@example.com"])
    ws.append(["NV002", "Trần Thị B", "SITE", "Kỹ sư", "SITE-A", "15/03/2024", "0909123457", "b.tran@example.com"])

    for col in ws.columns:
        max_len = max(len(str(c.value or "")) for c in col) + 2
        ws.column_dimensions[col[0].column_letter].width = max(max_len, 12)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=import_employee_template.xlsx"},
    )


@router.post("/import", status_code=200)
@limiter.limit("5/minute")
async def import_employees_from_excel(
    request: Request,
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
    perms: dict = Depends(RequirePermission('can_manage_employees')),
):
    if not file.filename or not file.filename.endswith(('.xlsx', '.xls')):
        logger.warning("Import rejected - invalid file type: %s by user %s",
                       file.filename, current_user.username)
        raise BadRequestError(detail=RESPONSE_MESSAGES["ONLY_EXCEL_SUPPORTED"])

    content = await file.read()
    if not content:
        logger.warning("Import rejected - empty file: %s by user %s",
                       file.filename, current_user.username)
        raise BadRequestError(detail=RESPONSE_MESSAGES["EMPTY_FILE"])

    file_size_kb = len(content) / 1024
    logger.info("Import started: file=%s size=%.1fKB user=%s",
                file.filename, file_size_kb, current_user.username)

    try:
        records, parse_errors = parse_excel(content)
    except ValueError as e:
        logger.warning("Parse failed for file=%s: %s", file.filename, e)
        raise BadRequestError(detail=str(e))

    if parse_errors:
        logger.info("Import validation failed: file=%s errors=%d",
                    file.filename, len(parse_errors))
        return {"errors": parse_errors, "imported": []}

    result = await import_employees(db, records)
    success = result.get("imported_count", result.get("imported", 0))
    if isinstance(success, list):
        success = len(success)
    logger.info("Import completed: file=%s imported=%d errors=%d user=%s",
                file.filename, success, len(result.get("errors", [])), current_user.username)
    await create_audit_log(
        db, action="employee.import", resource_type="employee",
        user_id=str(current_user.id),
        new_values={"file": file.filename, "imported": success, "errors": len(result.get("errors", []))},
        request=request, success=True,
    )
    return result


@router.get("/export")
async def export_employees_to_excel(
    department_id: str | None = Query(None),
    active_only: bool = Query(True),
    employee_ids: list[str] | None = Query(None),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
    perms: dict = Depends(get_current_permissions),
):
    query = select(Employee).options(
        selectinload(Employee.department),
        selectinload(Employee.workplace),
        selectinload(Employee.position_rel),
    )
    if active_only:
        query = query.where(Employee.is_active == True)

    scope_filter = get_scope_filter(perms, current_user)
    if scope_filter["employee_id"]:
        query = query.where(Employee.id == scope_filter["employee_id"])
    elif scope_filter["department_id"]:
        query = query.where(Employee.department_id == scope_filter["department_id"])
    elif department_id:
        query = query.where(Employee.department_id == department_id)

    if employee_ids:
        query = query.where(Employee.id.in_(employee_ids))
    query = query.order_by(Employee.full_name)

    result = await db.execute(query)
    employees = result.scalars().all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Nhân viên"

    headers = ["Mã NV", "Họ tên", "Phòng ban", "Chức vụ", "Nơi làm việc", "Ngày vào làm", "SĐT", "Email", "Trạng thái"]
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    for row_idx, emp in enumerate(employees, 2):
        row_data = [
            emp.employee_code or "",
            emp.full_name,
            emp.department.name if emp.department else "",
            emp.position_rel.name if emp.position_rel else (emp.position or ""),
            emp.workplace.name if emp.workplace else "",
            emp.hire_date.strftime("%d/%m/%Y") if emp.hire_date else "",
            emp.phone or "",
            emp.email or "",
            "Đang hoạt động" if emp.is_active else "Ngừng hoạt động",
        ]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border
            if col_idx in (1, 6, 9):
                cell.alignment = Alignment(horizontal="center")

    for col in ws.columns:
        max_length = 0
        for cell in col:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except (AttributeError, TypeError, ValueError):
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[col[0].column_letter].width = adjusted_width

    ws.auto_filter.ref = ws.dimensions

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=nhanvien_{timestamp}.xlsx"},
    )
